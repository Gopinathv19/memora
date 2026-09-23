"""Document Processor: local parsing and per-page triage. No model calls."""

import io

import pytest

from app.core.config import Settings
from app.processing import DocumentProcessor, UnsupportedDocumentError
from app.processing.triage import PageSignals, classify, junk_ratio
from app.schemas.enums import ExtractionMode, ExtractionRoute, PageDifficulty
from tests.extraction_fakes import LOREM, make_pdf, png_bytes, scanned_pdf, table_page, text_page

PDF = "application/pdf"


@pytest.fixture
def processor():
    return DocumentProcessor(Settings(database_url="postgresql+psycopg://unused/unused"))


def test_plain_text_page_is_easy_and_needs_no_model(processor):
    doc = processor.process(make_pdf([text_page(LOREM)]), PDF, "a.pdf")
    [unit] = doc.units
    assert unit.difficulty == PageDifficulty.EASY
    assert unit.route == ExtractionRoute.TEXT
    assert "INV-2041" in unit.text
    assert unit.page_image is None and unit.images == []


def test_table_page_is_hard_and_rendered_for_the_layout_model(processor):
    doc = processor.process(make_pdf([table_page()]), PDF, "t.pdf")
    [unit] = doc.units
    assert unit.signals["table_count"] >= 1
    assert unit.difficulty == PageDifficulty.HARD
    assert unit.route == ExtractionRoute.LAYOUT
    assert unit.page_image.mime == "image/png"

    from PIL import Image

    rendered = Image.open(io.BytesIO(unit.page_image.data))
    # Long side at the configured size, inside Nemotron-Parse's window.
    assert abs(max(rendered.size) - 1800) <= 1  # pdfium rounds up
    assert min(rendered.size) >= 1024


def test_scanned_page_without_text_is_hard(processor):
    doc = processor.process(scanned_pdf(), PDF, "scan.pdf")
    [unit] = doc.units
    assert unit.signals["text_chars"] == 0
    assert unit.difficulty == PageDifficulty.HARD
    assert unit.page_image is not None


def test_mixed_pdf_is_triaged_page_by_page(processor):
    data = make_pdf([text_page(LOREM), table_page(), text_page(LOREM)])
    doc = processor.process(data, PDF, "mixed.pdf")
    assert [u.difficulty for u in doc.units] == [
        PageDifficulty.EASY,
        PageDifficulty.HARD,
        PageDifficulty.EASY,
    ]
    assert [u.index for u in doc.units] == [1, 2, 3]


def test_deep_mode_sends_every_page_to_layout(processor):
    doc = processor.process(
        make_pdf([text_page(LOREM), text_page(LOREM)]), PDF, "a.pdf", ExtractionMode.DEEP
    )
    assert {u.route for u in doc.units} == {ExtractionRoute.LAYOUT}
    assert all(u.page_image is not None for u in doc.units)


def test_page_cap_is_recorded_not_silent():
    settings = Settings(database_url="postgresql+psycopg://x/y", extraction_max_pages=2)
    doc = DocumentProcessor(settings).process(
        make_pdf([text_page(LOREM)] * 4), PDF, "long.pdf"
    )
    assert len(doc.units) == 2
    assert doc.skipped_units == 2
    assert any("first 2 of 4" in w for w in doc.warnings)


def test_docx_keeps_tables_in_place_and_routes_pictures_to_vision(processor):
    import docx

    document = docx.Document()
    document.add_heading("Policy Schedule", level=1)
    document.add_paragraph("Policy number PN-778 issued to Arun Kumar.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Premium", "Due"
    table.cell(1, 0).text, table.cell(1, 1).text = "4,200", "2026-10-01"
    document.add_paragraph("Signed below.")
    document.add_picture(io.BytesIO(png_bytes()))
    out = io.BytesIO()
    document.save(out)

    doc = processor.process(out.getvalue(), None, "policy.docx")
    [unit] = doc.units
    assert unit.text.startswith("# Policy Schedule")
    assert "| Premium | Due |" in unit.text
    # Order preserved: the table sits between the two paragraphs.
    assert unit.text.index("PN-778") < unit.text.index("Premium") < unit.text.index("Signed")
    assert unit.route == ExtractionRoute.VISION
    assert len(unit.images) == 1


def test_pptx_is_one_unit_per_slide(processor):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Quarterly review"
    first.placeholders[1].text = "Revenue grew 18 percent"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(1))
    out = io.BytesIO()
    deck.save(out)

    doc = processor.process(out.getvalue(), None, "deck.pptx")
    assert [u.kind for u in doc.units] == ["slide", "slide"]
    assert "Revenue grew" in doc.units[0].text
    assert doc.units[0].route == ExtractionRoute.TEXT
    assert doc.units[1].route == ExtractionRoute.VISION
    assert len(doc.units[1].images) == 1


def test_xlsx_sheets_become_markdown_tables(processor):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Claims"
    sheet.append(["Claim", "Amount"])
    sheet.append(["CL-1", 900])
    out = io.BytesIO()
    book.save(out)

    doc = processor.process(out.getvalue(), None, "claims.xlsx")
    [unit] = doc.units
    assert unit.kind == "sheet"
    assert "## Sheet: Claims" in unit.text
    assert "| CL-1 | 900 |" in unit.text
    assert unit.route == ExtractionRoute.TEXT


def test_image_file_goes_to_vision(processor):
    doc = processor.process(png_bytes(), "image/png", "photo.png")
    [unit] = doc.units
    assert unit.kind == "image"
    assert unit.route == ExtractionRoute.VISION
    assert unit.images[0].mime == "image/png"


def test_text_file_is_read_directly(processor):
    doc = processor.process("Name: Arun\nCity: Chennai".encode(), "text/plain", "n.txt")
    assert doc.units[0].text == "Name: Arun\nCity: Chennai"


@pytest.mark.parametrize(
    ("mime", "filename"),
    [
        ("application/zip", "archive.zip"),
        ("application/msword", "old.doc"),
        (None, None),
    ],
)
def test_unsupported_types_are_rejected_clearly(processor, mime, filename):
    with pytest.raises(UnsupportedDocumentError):
        processor.process(b"whatever", mime, filename)


def test_unreadable_image_is_rejected(processor):
    with pytest.raises(UnsupportedDocumentError):
        processor.process(b"not really a png", "image/png", "x.png")


def _signals(**overrides):
    base = dict(
        text_chars=1500, image_count=0, image_ratio=0.0,
        table_count=0, ruled_lines=0, junk_ratio=0.0,
    )
    return PageSignals(**{**base, **overrides})


@pytest.mark.parametrize(
    ("overrides", "expected"),
    [
        ({}, PageDifficulty.EASY),
        ({"image_count": 1, "image_ratio": 0.05}, PageDifficulty.MEDIUM),
        ({"image_count": 1, "image_ratio": 0.4}, PageDifficulty.HARD),
        ({"table_count": 1}, PageDifficulty.HARD),
        ({"ruled_lines": 12}, PageDifficulty.HARD),
        ({"junk_ratio": 0.3}, PageDifficulty.HARD),
        ({"text_chars": 0, "image_count": 1, "image_ratio": 0.1}, PageDifficulty.HARD),
        ({"image_count": 5, "image_ratio": 0.1}, PageDifficulty.HARD),
        ({"text_chars": 0}, PageDifficulty.EASY),  # a blank page
    ],
)
def test_triage_rules(overrides, expected):
    settings = Settings(database_url="postgresql+psycopg://x/y")
    assert classify(_signals(**overrides), settings) == expected


def test_junk_ratio_spots_unmapped_glyphs():
    assert junk_ratio("clean text") == 0
    assert junk_ratio("(cid:12)(cid:40)(cid:7)ab") > 0.8
